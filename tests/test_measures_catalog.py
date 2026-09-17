"""Keep dax/measures.dax aligned with Excel _Measures documentation sheet."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Northwind_DW_DimFact.xlsx"
MEASURES_DAX = ROOT / "dax" / "measures.dax"


def _normalize_dax(expr: str) -> str:
    """Collapse whitespace for stable comparison."""
    return re.sub(r"\s+", " ", expr.strip())


def parse_measures_dax(path: Path) -> dict[str, str]:
    """Parse measure names and expression bodies from measures.dax."""
    measures: dict[str, str] = {}
    current_name: str | None = None
    body_lines: list[str] = []

    for raw in path.read_text(encoding="utf-8").splitlines():
        if raw.strip().startswith("//"):
            continue
        is_measure_header = (
            raw
            and not raw[0].isspace()
            and "=" in raw
            and not raw.lstrip().startswith(("VAR ", "RETURN"))
        )
        if is_measure_header:
            if current_name is not None:
                measures[current_name] = _normalize_dax(" ".join(body_lines))
            name_part, _, rest = raw.partition("=")
            current_name = name_part.strip()
            body_lines = [rest.strip()] if rest.strip() else []
        elif current_name is not None:
            body_lines.append(raw.strip())

    if current_name is not None:
        measures[current_name] = _normalize_dax(" ".join(body_lines))
    return measures


def load_excel_measures(path: Path) -> dict[str, str]:
    """Load measure catalog from _Measures sheet (Measure -> DAX)."""
    wb = load_workbook(path, data_only=True)
    ws = wb["_Measures"]
    out: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        out[str(row[1])] = _normalize_dax(str(row[3] or ""))
    wb.close()
    return out


@pytest.fixture(scope="module")
def dax_catalog() -> dict[str, str]:
    assert MEASURES_DAX.exists(), f"Missing {MEASURES_DAX}"
    return parse_measures_dax(MEASURES_DAX)


@pytest.fixture(scope="module")
def excel_catalog() -> dict[str, str]:
    assert XLSX.exists(), f"Missing {XLSX}"
    return load_excel_measures(XLSX)


def test_measures_dax_parses(dax_catalog: dict[str, str]) -> None:
    assert len(dax_catalog) >= 24
    assert "Sales Amount" in dax_catalog
    assert dax_catalog["Sales Amount"] == "SUM ( FactSales[SalesAmount] )"


def test_excel_measures_sheet_not_empty(excel_catalog: dict[str, str]) -> None:
    assert len(excel_catalog) >= 24


def test_dax_catalog_matches_excel(dax_catalog: dict[str, str], excel_catalog: dict[str, str]) -> None:
    missing_in_excel = set(dax_catalog) - set(excel_catalog)
    missing_in_dax = set(excel_catalog) - set(dax_catalog)
    assert not missing_in_excel, f"Add to _Measures via extend_northwind_dw.py: {sorted(missing_in_excel)}"
    assert not missing_in_dax, f"Remove or add to measures.dax: {sorted(missing_in_dax)}"

    mismatches = []
    for name, dax_body in dax_catalog.items():
        excel_body = excel_catalog[name]
        if dax_body != excel_body:
            mismatches.append((name, dax_body, excel_body))
    assert not mismatches, mismatches[:3]
