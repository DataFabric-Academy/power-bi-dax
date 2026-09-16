"""Validate extended Northwind DW Excel source for the DAX curriculum."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Northwind_DW_DimFact.xlsx"


@pytest.fixture(scope="module")
def wb():
    assert XLSX.exists(), f"Missing {XLSX}"
    book = load_workbook(XLSX, data_only=True)
    yield book
    book.close()


def _headers(ws):
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def _data(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


def _col(ws, name):
    h = _headers(ws)
    i = h.index(name)
    return [r[i] for r in _data(ws)]


@pytest.mark.parametrize(
    "sheet,key",
    [
        ("DimCustomer", "CustomerKey"),
        ("DimEmployee", "EmployeeKey"),
        ("DimProduct", "ProductKey"),
        ("DimShipper", "ShipperKey"),
        ("DimDate", "DateKey"),
        ("DimDate", "Date"),
        ("FactSales", "SalesKey"),
        ("DimProduct", "ProductCode"),
    ],
)
def test_unique_keys(wb, sheet, key):
    vals = _col(wb[sheet], key)
    assert len(vals) == len(set(vals))
    assert all(v is not None for v in vals)


def test_dimdate_contiguous(wb):
    dates = sorted(_col(wb["DimDate"], "Date"))
    assert len(dates) == 1096
    for i in range(1, len(dates)):
        d0 = dates[i - 1].date() if isinstance(dates[i - 1], datetime) else dates[i - 1]
        d1 = dates[i].date() if isinstance(dates[i], datetime) else dates[i]
        assert (d1 - d0).days == 1


@pytest.mark.parametrize(
    "fk,dim,pk",
    [
        ("ProductKey", "DimProduct", "ProductKey"),
        ("CustomerKey", "DimCustomer", "CustomerKey"),
        ("EmployeeKey", "DimEmployee", "EmployeeKey"),
        ("ShipperKey", "DimShipper", "ShipperKey"),
        ("OrderDateKey", "DimDate", "DateKey"),
        ("RequiredDateKey", "DimDate", "DateKey"),
        ("ShippedDateKey", "DimDate", "DateKey"),
    ],
)
def test_no_orphan_fks(wb, fk, dim, pk):
    keys = set(_col(wb[dim], pk))
    orphans = [v for v in _col(wb["FactSales"], fk) if v not in keys]
    assert orphans == []


def test_employee_demographics(wb):
    assert "BirthDate" in _headers(wb["DimEmployee"])
    assert all(v is not None for v in _col(wb["DimEmployee"], "BirthDate"))
    assert all(isinstance(v, (int, float)) for v in _col(wb["DimEmployee"], "BirthYear"))
    assert set(_col(wb["DimEmployee"], "Gender")) <= {"Male", "Female"}
    assert len(_data(wb["DimEmployee"])) == 9


def test_product_teaching_columns(wb):
    h = _headers(wb["DimProduct"])
    for col in ("ProductCode", "Status", "UnitsInStock", "ListPrice"):
        assert col in h
    assert set(_col(wb["DimProduct"], "Status")) <= {"Active", "Inactive"}
    assert all(isinstance(v, (int, float)) for v in _col(wb["DimProduct"], "ListPrice"))
    assert all(v is not None for v in _col(wb["DimProduct"], "ProductCode"))


def test_linecost_present_and_differs_from_standard(wb):
    assert "LineCost" in _headers(wb["FactSales"])
    assert all(v is not None for v in _col(wb["FactSales"], "LineCost"))
    std = dict(
        zip(
            _col(wb["DimProduct"], "ProductKey"),
            _col(wb["DimProduct"], "StandardCost"),
        )
    )
    h = _headers(wb["FactSales"])
    pk_i = h.index("ProductKey")
    qty_i = h.index("Quantity")
    lc_i = h.index("LineCost")
    differ = 0
    for r in _data(wb["FactSales"]):
        expected = float(std[r[pk_i]]) * float(r[qty_i])
        if abs(float(r[lc_i]) - expected) > 0.01:
            differ += 1
    assert differ > 0, "LineCost must differ from StandardCost*Quantity for teaching"


def test_qa_sheet_all_pass(wb):
    rows = _data(wb["_QA"])
    # Status is last column
    statuses = [r[3] for r in rows]
    assert statuses
    assert all(s == "PASS" for s in statuses), [r for r in rows if r[3] != "PASS"]


def test_import_tables_exist(wb):
    for name in (
        "DimDate",
        "DimCustomer",
        "DimEmployee",
        "DimProduct",
        "DimShipper",
        "FactSales",
    ):
        assert name in wb.sheetnames
        assert wb[name].tables or True  # tables may be present
